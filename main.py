from __future__ import annotations

import base64
import hashlib
import hmac
import json
import os
import secrets
import threading

from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import (
    FastAPI,
    Depends,
    HTTPException,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, EmailStr, Field


# =========================================================
# CONFIGURATION
# =========================================================

BASE = Path(__file__).resolve().parent

DATA_DIR = BASE / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)

XLSX_PATH = Path(
    os.getenv(
        "XLSX_PATH",
        str(DATA_DIR / "college_data.xlsx"),
    )
)

SECRET_KEY = os.getenv(
    "ASHISH_SECRET_KEY",
    "DEV_ONLY_CHANGE_THIS_SECRET_BEFORE_DEPLOYMENT",
)

if SECRET_KEY == "DEV_ONLY_CHANGE_THIS_SECRET_BEFORE_DEPLOYMENT":
    print(
        "WARNING: Development token secret in use. "
        "Set ASHISH_SECRET_KEY before deployment."
    )

app = FastAPI(
    title="Ashish College Connect API",
    version="1.0.0",
    description="College portal API with Excel workbook storage.",
)

# Swagger will use this security scheme for protected endpoints.
bearer_scheme = HTTPBearer(auto_error=False)

origins = [
    item.strip()
    for item in os.getenv("CORS_ORIGINS", "*").split(",")
    if item.strip()
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

lock = threading.RLock()


# =========================================================
# EXCEL WORKBOOK CONFIGURATION
# =========================================================

HEADERS = {
    "Users": [
        "id",
        "name",
        "email",
        "password_hash",
        "role",
        "phone",
    ],
    "FAQs": [
        "id",
        "category",
        "question",
        "answer",
    ],
    "Queries": [
        "id",
        "name",
        "email",
        "phone",
        "message",
        "status",
        "created_at",
    ],
    "Announcements": [
        "id",
        "title",
        "content",
        "created_at",
    ],
}

SEED_FAQS = [
    (
        "Admissions",
        "How can I apply for admission?",
        "Visit the college admissions office or official admissions "
        "portal and follow the current application instructions.",
    ),
    (
        "Admissions",
        "Which documents are generally required for admission?",
        "Commonly requested documents include qualifying exam marks "
        "cards, transfer certificate, identity proof, and photographs. "
        "Confirm the current list with the admissions office.",
    ),
    (
        "Fees",
        "How can I pay my college fees?",
        "Contact the accounts section for current payment methods, "
        "due dates, and receipts.",
    ),
    (
        "Examinations",
        "Where can I find exam timetables?",
        "Check official college announcements and the examination "
        "section for the latest timetable.",
    ),
    (
        "Examinations",
        "How do I get my hall ticket?",
        "Follow the examination section instructions and contact "
        "them if your hall ticket is unavailable.",
    ),
    (
        "Library",
        "What are the library timings?",
        "Check the latest library notice or contact the library desk "
        "for current opening hours.",
    ),
    (
        "Placements",
        "How can I contact the placement cell?",
        "Contact the college placement cell through the official "
        "college directory or student portal.",
    ),
    (
        "Certificates",
        "How can I request a bonafide certificate?",
        "Contact the administrative office and ask about the request "
        "form, required documents, and processing time.",
    ),
    (
        "Hostel",
        "How do I enquire about hostel accommodation?",
        "Contact the hostel office for availability, fees, rules, "
        "and application details.",
    ),
    (
        "General",
        "How can I contact the college?",
        "Submit a query through this portal or contact the college "
        "administrative office.",
    ),
]


# =========================================================
# GENERAL HELPERS
# =========================================================

def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def hash_password(password: str, salt: Optional[str] = None) -> str:
    salt = salt or secrets.token_hex(16)

    hashed = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode(),
        salt.encode(),
        200_000,
    ).hex()

    return salt + "$" + hashed


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, expected = stored.split("$", 1)

        actual = hashlib.pbkdf2_hmac(
            "sha256",
            password.encode(),
            salt.encode(),
            200_000,
        ).hex()

        return hmac.compare_digest(actual, expected)

    except Exception:
        return False


# =========================================================
# EXCEL WORKBOOK HELPERS
# =========================================================

def ensure_workbook():
    with lock:
        if XLSX_PATH.exists():
            wb = load_workbook(XLSX_PATH)
        else:
            wb = Workbook()
            wb.remove(wb.active)

        for sheet_name, required_headers in HEADERS.items():
            if sheet_name not in wb.sheetnames:
                ws = wb.create_sheet(sheet_name)
                ws.append(required_headers)

            else:
                ws = wb[sheet_name]

                # Create headers if the sheet is completely empty.
                if ws.max_row == 0 or (
                    ws.max_row == 1
                    and all(cell.value is None for cell in ws[1])
                ):
                    ws.append(required_headers)

                # Add missing columns at the end without deleting
                # existing workbook data.
                elif ws.max_row >= 1:
                    existing_headers = [
                        cell.value for cell in ws[1]
                    ]

                    for header in required_headers:
                        if header not in existing_headers:
                            ws.cell(
                                row=1,
                                column=len(existing_headers) + 1,
                                value=header,
                            )
                            existing_headers.append(header)

        # Seed FAQs if the FAQ sheet contains only its header row.
        if wb["FAQs"].max_row <= 1:
            for index, faq in enumerate(SEED_FAQS, start=1):
                wb["FAQs"].append([index, *faq])

        # Add demo users if their emails do not already exist.
        ws = wb["Users"]

        user_headers = [
            cell.value for cell in ws[1]
        ]

        email_column = user_headers.index("email") + 1

        existing_emails = {
            str(ws.cell(row=row, column=email_column).value).lower()
            for row in range(2, ws.max_row + 1)
            if ws.cell(row=row, column=email_column).value
        }

        if "admin@college.edu" not in existing_emails:
            ws.append([
                1,
                "College Admin",
                "admin@college.edu",
                hash_password("Admin@123"),
                "admin",
                "",
            ])

        if "student@college.edu" not in existing_emails:
            ws.append([
                2,
                "Demo Student",
                "student@college.edu",
                hash_password("Student@123"),
                "student",
                "",
            ])

        wb.save(XLSX_PATH)
        wb.close()


def rows(sheet_name: str) -> list[dict]:
    ensure_workbook()

    with lock:
        wb = load_workbook(
            XLSX_PATH,
            read_only=True,
            data_only=True,
        )

        values = list(
            wb[sheet_name].iter_rows(values_only=True)
        )

        wb.close()

    if not values:
        return []

    headers = [str(value) for value in values[0]]

    return [
        {
            headers[index]: (
                row[index] if index < len(row) else None
            )
            for index in range(len(headers))
        }
        for row in values[1:]
        if any(value is not None for value in row)
    ]


def append_row(sheet_name: str, values: list):
    ensure_workbook()

    with lock:
        wb = load_workbook(XLSX_PATH)
        ws = wb[sheet_name]

        headers = [
            cell.value for cell in ws[1]
        ]

        # Map the supplied values to the actual column order.
        row_data = dict(zip(HEADERS[sheet_name], values))

        ws.append([
            row_data.get(header)
            for header in headers
        ])

        wb.save(XLSX_PATH)
        wb.close()


def next_id(sheet_name: str) -> int:
    existing_ids = [
        int(item.get("id") or 0)
        for item in rows(sheet_name)
    ]

    return max(existing_ids or [0]) + 1


# =========================================================
# TOKEN AUTHENTICATION
# =========================================================

def make_token(email: str, role: str) -> str:
    payload = json.dumps(
        {
            "email": email,
            "role": role,
            "exp": int(
                datetime.now(timezone.utc).timestamp()
            ) + 28800,
        },
        separators=(",", ":"),
    ).encode()

    body = (
        base64.urlsafe_b64encode(payload)
        .decode()
        .rstrip("=")
    )

    signature = hmac.new(
        SECRET_KEY.encode(),
        body.encode(),
        hashlib.sha256,
    ).hexdigest()

    return body + "." + signature


def decode_token(token: str) -> dict:
    try:
        body, signature = token.split(".", 1)

        expected_signature = hmac.new(
            SECRET_KEY.encode(),
            body.encode(),
            hashlib.sha256,
        ).hexdigest()

        if not hmac.compare_digest(
            signature,
            expected_signature,
        ):
            raise ValueError("Invalid signature")

        payload = json.loads(
            base64.urlsafe_b64decode(
                body + "=" * (-len(body) % 4)
            )
        )

        if int(payload["exp"]) <= int(
            datetime.now(timezone.utc).timestamp()
        ):
            raise ValueError("Expired token")

        return payload

    except Exception:
        raise HTTPException(
            status_code=401,
            detail="Invalid or expired token. Please log in again.",
        )


# =========================================================
# FIXED AUTHENTICATION DEPENDENCIES
# =========================================================

def current_user(
    credentials: Optional[
        HTTPAuthorizationCredentials
    ] = Depends(bearer_scheme),
):
    if credentials is None:
        raise HTTPException(
            status_code=401,
            detail="Please log in.",
        )

    if credentials.scheme.lower() != "bearer":
        raise HTTPException(
            status_code=401,
            detail="Please log in using a Bearer token.",
        )

    return decode_token(credentials.credentials)


def admin_user(user=Depends(current_user)):
    if user.get("role") != "admin":
        raise HTTPException(
            status_code=403,
            detail="Admin access required.",
        )

    return user


# =========================================================
# REQUEST MODELS
# =========================================================

class LoginRequest(BaseModel):
    email: EmailStr
    password: str = Field(min_length=1, max_length=128)


class QueryRequest(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    email: EmailStr
    message: str = Field(min_length=1, max_length=3000)
    phone: str = Field(default="", max_length=30)


class FAQRequest(BaseModel):
    category: str = Field(min_length=1, max_length=80)
    question: str = Field(min_length=1, max_length=500)
    answer: str = Field(min_length=1, max_length=3000)


class AnnouncementRequest(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    content: str = Field(min_length=1, max_length=5000)


class StatusRequest(BaseModel):
    status: str = Field(
        pattern="^(open|in_progress|resolved)$"
    )


# =========================================================
# STARTUP
# =========================================================

@app.on_event("startup")
def startup():
    ensure_workbook()


# =========================================================
# PUBLIC ENDPOINTS
# =========================================================

@app.get("/")
def root():
    return {
        "app": "Ashish College Connect API",
        "status": "running",
        "docs": "/docs",
    }


@app.get("/health")
def health():
    ensure_workbook()

    return {
        "status": "ok",
        "excel_file_exists": XLSX_PATH.exists(),
        "excel_file": XLSX_PATH.name,
    }


@app.get("/api/faqs")
def get_faqs():
    return {
        "items": rows("FAQs"),
    }


@app.get("/api/announcements")
def get_announcements():
    return {
        "items": list(reversed(rows("Announcements"))),
    }


# =========================================================
# LOGIN AND CURRENT USER
# =========================================================

@app.post("/api/login")
def login(body: LoginRequest):
    email = str(body.email).lower()

    for user in rows("Users"):
        if (
            str(user.get("email", "")).lower() == email
            and verify_password(
                body.password,
                str(user.get("password_hash", "")),
            )
        ):
            return {
                "access_token": make_token(
                    email,
                    str(user.get("role", "student")),
                ),
                "token_type": "bearer",
                "user": {
                    "id": user["id"],
                    "name": user["name"],
                    "email": email,
                    "role": user["role"],
                },
            }

    raise HTTPException(
        status_code=401,
        detail="Incorrect email or password.",
    )


@app.get("/api/me")
def me(user=Depends(current_user)):
    matches = [
        item for item in rows("Users")
        if str(item.get("email", "")).lower()
        == user["email"].lower()
    ]

    if not matches:
        raise HTTPException(
            status_code=401,
            detail="User not found.",
        )

    item = matches[0]

    return {
        key: item.get(key)
        for key in ("id", "name", "email", "role")
    }


# =========================================================
# STUDENT QUERIES
# =========================================================

@app.post("/api/queries")
def create_query(body: QueryRequest):
    query_id = next_id("Queries")

    append_row(
        "Queries",
        [
            query_id,
            body.name.strip(),
            str(body.email).lower(),
            body.phone.strip(),
            body.message.strip(),
            "open",
            now(),
        ],
    )

    return {
        "message": "Your query has been submitted.",
        "query_id": query_id,
        "status": "open",
    }


@app.get("/api/queries")
def get_queries(user=Depends(current_user)):
    items = rows("Queries")

    if user.get("role") != "admin":
        items = [
            item for item in items
            if str(item.get("email", "")).lower()
            == user["email"].lower()
        ]

    return {
        "items": list(reversed(items)),
    }


# =========================================================
# ADMIN DASHBOARD
# =========================================================

@app.get("/api/admin/summary")
def admin_summary(_=Depends(admin_user)):
    users = rows("Users")
    queries = rows("Queries")
    faqs_data = rows("FAQs")
    announcements_data = rows("Announcements")

    statuses = [
        str(item.get("status", "open"))
        for item in queries
    ]

    return {
        "users": len(users),
        "faqs": len(faqs_data),
        "announcements": len(announcements_data),
        "queries": len(queries),
        "open_queries": statuses.count("open"),
        "in_progress_queries": statuses.count("in_progress"),
        "resolved_queries": statuses.count("resolved"),
    }


@app.get("/api/admin/users")
def admin_users(_=Depends(admin_user)):
    return {
        "items": [
            {
                key: item.get(key)
                for key in (
                    "id",
                    "name",
                    "email",
                    "role",
                    "phone",
                )
            }
            for item in rows("Users")
        ]
    }


@app.get("/api/admin/queries")
def admin_queries(_=Depends(admin_user)):
    return {
        "items": list(reversed(rows("Queries"))),
    }


@app.patch("/api/admin/queries/{query_id}/status")
def update_status(
    query_id: int,
    body: StatusRequest,
    _=Depends(admin_user),
):
    ensure_workbook()
    found = False

    with lock:
        wb = load_workbook(XLSX_PATH)
        ws = wb["Queries"]

        headers = [
            cell.value for cell in ws[1]
        ]

        id_column = headers.index("id") + 1
        status_column = headers.index("status") + 1

        for row_number in range(2, ws.max_row + 1):
            cell_id = ws.cell(
                row=row_number,
                column=id_column,
            ).value

            if cell_id == query_id:
                ws.cell(
                    row=row_number,
                    column=status_column,
                    value=body.status,
                )
                found = True
                break

        if found:
            wb.save(XLSX_PATH)

        wb.close()

    if not found:
        raise HTTPException(
            status_code=404,
            detail="Query not found.",
        )

    return {
        "message": "Query status updated.",
        "query_id": query_id,
        "status": body.status,
    }


# =========================================================
# ADMIN FAQ MANAGEMENT
# =========================================================

@app.post("/api/admin/faqs")
def add_faq(
    body: FAQRequest,
    _=Depends(admin_user),
):
    faq_id = next_id("FAQs")

    append_row(
        "FAQs",
        [
            faq_id,
            body.category.strip(),
            body.question.strip(),
            body.answer.strip(),
        ],
    )

    return {
        "message": "FAQ added.",
        "id": faq_id,
    }


# =========================================================
# ADMIN ANNOUNCEMENT MANAGEMENT
# =========================================================

@app.post("/api/admin/announcements")
def add_announcement(
    body: AnnouncementRequest,
    _=Depends(admin_user),
):
    announcement_id = next_id("Announcements")

    append_row(
        "Announcements",
        [
            announcement_id,
            body.title.strip(),
            body.content.strip(),
            now(),
        ],
    )

    return {
        "message": "Announcement published.",
        "id": announcement_id,
    }